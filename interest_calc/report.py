"""
Moduł do generowania raportów w formatach CSV i XLSX.

Obsługuje:
- Eksport wyników do CSV
- Eksport wyników do XLSX z formatowaniem
- Generowanie timestampów dla nazw plików
"""

from datetime import datetime
from pathlib import Path
from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

from interest_calc.models import CalculationResult


def generate_report_filename(prefix: str = "interest_report", extension: str = "csv") -> str:
    """
    Generuje nazwę pliku raportu z timestampem.

    Args:
        prefix: Prefiks nazwy pliku
        extension: Rozszerzenie pliku (bez kropki)

    Returns:
        Nazwa pliku w formacie: prefix_YYYYMMDD_HHMMSS.extension
    """
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"{prefix}_{timestamp}.{extension}"


def create_summary_dict(result: CalculationResult) -> dict:
    """
    Tworzy słownik z podsumowaniem obliczeń.

    Args:
        result: Wynik obliczeń

    Returns:
        Słownik z kluczowymi informacjami
    """
    summary = {
        "Kwota główna (PLN)": result.settings.principal_pln,
        "Data początkowa": result.settings.start_date.isoformat(),
        "Data końcowa": result.settings.end_date.isoformat(),
        "Liczba dni": result.total_days,
        "Średnia ważona stopa (%)": round(result.weighted_avg_rate, 4),
        "Łączne odsetki (PLN)": round(result.total_interest, 2),
        "Łączne roszczenie (PLN)": round(
            result.settings.principal_pln + result.total_interest, 2
        ),
        "Podstawa liczenia dni": result.settings.basis.value,
        "Liczba okresów stawek": len(result.segments),
    }

    if result.settings.principal_chf:
        summary["Kwota główna (CHF)"] = result.settings.principal_chf

    return summary


def create_segments_dataframe(result: CalculationResult) -> pd.DataFrame:
    """
    Tworzy DataFrame z segmentami obliczeń.

    Args:
        result: Wynik obliczeń

    Returns:
        DataFrame z kolumnami: Data od, Data do, Dni, Stopa (%), Podstawa (PLN), Odsetki (PLN)
    """
    data = []

    for segment in result.segments:
        data.append(
            {
                "Data od": segment.start_date.isoformat(),
                "Data do": segment.end_date.isoformat(),
                "Dni": segment.days,
                "Stopa (%)": round(segment.rate_percent, 2),
                "Podstawa (PLN)": round(segment.principal, 2),
                "Odsetki (PLN)": round(segment.interest, 2),
            }
        )

    df = pd.DataFrame(data)

    # Dodaj wiersz z sumami
    if not df.empty:
        totals = {
            "Data od": "SUMA",
            "Data do": "",
            "Dni": df["Dni"].sum(),
            "Stopa (%)": "",  # Średnią mamy już w podsumowaniu
            "Podstawa (PLN)": "",  # Nie sumujemy podstaw (mogą się powtarzać)
            "Odsetki (PLN)": round(df["Odsetki (PLN)"].sum(), 2),
        }
        df = pd.concat([df, pd.DataFrame([totals])], ignore_index=True)

    return df


def save_report_csv(result: CalculationResult, output_dir: Path) -> Path:
    """
    Zapisuje raport do pliku CSV.

    Args:
        result: Wynik obliczeń
        output_dir: Katalog wyjściowy

    Returns:
        Ścieżka do zapisanego pliku

    Raises:
        IOError: Jeśli nie można zapisać pliku
    """
    output_dir.mkdir(parents=True, exist_ok=True)

    filename = generate_report_filename(extension="csv")
    filepath = output_dir / filename

    try:
        # Utwórz DataFrame z podsumowaniem
        summary_dict = create_summary_dict(result)
        summary_df = pd.DataFrame(
            [{"Parametr": k, "Wartość": v} for k, v in summary_dict.items()]
        )

        # Utwórz DataFrame z segmentami
        segments_df = create_segments_dataframe(result)

        # Zapisz oba do jednego pliku CSV (rozdzielone pustą linią)
        with open(filepath, "w", encoding="utf-8") as f:
            f.write("PODSUMOWANIE OBLICZEŃ ODSETEK USTAWOWYCH ZA OPÓŹNIENIE\n")
            f.write("=" * 80 + "\n\n")
            summary_df.to_csv(f, index=False, encoding="utf-8")
            f.write("\n\nSZCZEGÓŁOWE SEGMENTY OBLICZEŃ\n")
            f.write("=" * 80 + "\n\n")
            segments_df.to_csv(f, index=False, encoding="utf-8")
            f.write("\n\n")
            f.write("UWAGA: Narzędzie nie stanowi porady prawnej.\n")
            f.write("Stawki należy zweryfikować i uaktualnić.\n")

        return filepath

    except Exception as e:
        raise IOError(f"Nie można zapisać pliku CSV {filepath}: {e}") from e


def save_report_xlsx(result: CalculationResult, output_dir: Path) -> Path:
    """
    Zapisuje raport do pliku XLSX z formatowaniem.

    Args:
        result: Wynik obliczeń
        output_dir: Katalog wyjściowy

    Returns:
        Ścieżka do zapisanego pliku

    Raises:
        IOError: Jeśli nie można zapisać pliku
    """
    output_dir.mkdir(parents=True, exist_ok=True)

    filename = generate_report_filename(extension="xlsx")
    filepath = output_dir / filename

    try:
        # Utwórz workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Raport odsetek"

        # Style
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        title_font = Font(bold=True, size=14)
        bold_font = Font(bold=True)
        center_align = Alignment(horizontal="center", vertical="center")

        # Tytuł
        ws["A1"] = "PODSUMOWANIE OBLICZEŃ ODSETEK USTAWOWYCH ZA OPÓŹNIENIE"
        ws["A1"].font = title_font
        ws.merge_cells("A1:B1")

        # Podsumowanie
        summary_dict = create_summary_dict(result)
        row = 3
        for key, value in summary_dict.items():
            ws[f"A{row}"] = key
            ws[f"B{row}"] = value
            ws[f"A{row}"].font = bold_font
            row += 1

        # Pusta linia
        row += 1

        # Nagłówek segmentów
        ws[f"A{row}"] = "SZCZEGÓŁOWE SEGMENTY OBLICZEŃ"
        ws[f"A{row}"].font = title_font
        ws.merge_cells(f"A{row}:F{row}")
        row += 2

        # Tabela segmentów
        segments_df = create_segments_dataframe(result)

        # Nagłówki tabeli
        header_row = row
        for col_num, column_name in enumerate(segments_df.columns, start=1):
            cell = ws.cell(row=header_row, column=col_num, value=column_name)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        # Dane
        for r_idx, row_data in enumerate(segments_df.values, start=header_row + 1):
            for c_idx, value in enumerate(row_data, start=1):
                ws.cell(row=r_idx, column=c_idx, value=value)

        # Pogrub ostatni wiersz (suma)
        if len(segments_df) > 0:
            last_row = header_row + len(segments_df)
            for col_num in range(1, len(segments_df.columns) + 1):
                ws.cell(row=last_row, column=col_num).font = bold_font

        # Dostosuj szerokości kolumn
        for col_num in range(1, ws.max_column + 1):
            max_length = 0
            column_letter = get_column_letter(col_num)
            for row_num in range(1, ws.max_row + 1):
                cell = ws.cell(row=row_num, column=col_num)
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        # Disclaimer na dole
        disclaimer_row = header_row + len(segments_df) + 2
        ws[f"A{disclaimer_row}"] = "⚠️ UWAGA: Narzędzie nie stanowi porady prawnej."
        ws[f"A{disclaimer_row + 1}"] = (
            "Stawki należy zweryfikować i uaktualnić. "
            "Skonsultuj się z prawnikiem w sprawie szczegółów Twojej sprawy."
        )

        # Zapisz
        wb.save(filepath)
        return filepath

    except Exception as e:
        raise IOError(f"Nie można zapisać pliku XLSX {filepath}: {e}") from e


def save_reports(result: CalculationResult, output_dir: Path) -> tuple[Path, Path]:
    """
    Zapisuje raporty w formatach CSV i XLSX.

    Args:
        result: Wynik obliczeń
        output_dir: Katalog wyjściowy

    Returns:
        Tuple (ścieżka_csv, ścieżka_xlsx)

    Raises:
        IOError: Jeśli nie można zapisać plików
    """
    csv_path = save_report_csv(result, output_dir)
    xlsx_path = save_report_xlsx(result, output_dir)

    return csv_path, xlsx_path
